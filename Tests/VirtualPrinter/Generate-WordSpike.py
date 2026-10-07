"""Generate a 120-page structural fixture, not the full acceptance W1 sample."""
import hashlib
import json
import sys
from pathlib import Path
from docx import Document
from docx.shared import Mm, Pt

out = Path(sys.argv[1])
out.mkdir(parents=True, exist_ok=True)
doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Mm(210), Mm(297)
section.top_margin = section.bottom_margin = Mm(20)
doc.styles['Normal'].font.name = 'Arial'
doc.styles['Normal'].font.size = Pt(11)
doc.core_properties.title = 'XT Phase 0 Word 120 Pages'
doc.core_properties.author = 'XT Virtual Printer Spike'
headings, pages = [], []
sentence = 'Tiếng Việt: Nguyễn Xuân Tới, đường bộ, cầu vượt, kỹ thuật, ă â ê ô ơ ư đ; á à ả ã ạ, ắ ằ ẳ ẵ ặ, ấ ầ ẩ ẫ ậ, ế ề ể ễ ệ, ố ồ ổỗ ộ, ớ ờ ở ỡ ợ, ứ ừử ữ ự.'
for page in range(1, 121):
    if page > 1:
        doc.add_page_break()
    offset = (page - 1) % 10
    chapter = (page - 1) // 10 + 1
    level = 1 if offset == 0 else 2 if offset in (1, 3, 5, 7) else 3 if offset in (2, 4, 6, 8) else 0
    title = f'Chương {chapter:02d}' if level == 1 else f'Mục {chapter:02d}.{offset:02d} cấp {level}'
    text = []
    if level:
        doc.add_heading(title, level)
        headings.append({'title': title, 'level': level, 'page': page})
        text.append(title)
    marker = f'PAGE-{page:03d}'
    doc.add_paragraph(marker)
    doc.add_paragraph(sentence)
    text.extend([marker, sentence])
    if page in (15, 45, 75):
        table = doc.add_table(rows=3, cols=3)
        table.style = 'Table Grid'
        for row in range(3):
            for col in range(3):
                value = f'Ô {row + 1}-{col + 1}'
                table.cell(row, col).text = value
                text.append(value)
    pages.append({'page': page, 'text': '\n'.join(text)})
path = out / 'word-120-structural.docx'
doc.save(path)
(out / 'word-expected.json').write_text(json.dumps({'pages': pages, 'headings': headings, 'sha256': hashlib.sha256(path.read_bytes()).hexdigest(), 'scope': 'Phase 0 structural fixture; full W1 TOC/images/spanning tables/mixed orientation are not included'}, ensure_ascii=False, indent=2), encoding='utf-8')
print(f'{path}: 120 intended pages; {len(headings)} headings')
